"""Turn whatever a county sends into a format the extractors read.

Counties answer records requests with whatever their office uses: Apple
Numbers, old Excel, OpenDocument, old Word, RTF, phone photos of paper
forms, zip archives of all of the above. Rather than teach every extractor
every format, this converts each file once, up front:

=====================  =====================================================
received               converted to
=====================  =====================================================
.numbers .ods .xls     .xlsx, one sheet per table (spreadsheet extractor)
.doc .rtf .odt .html   .txt (permit, form and enforcement extractors)
.jpg .png .tif .heic   .pdf, one page per image (read by OCR like any scan)
.zip                   its contents, each converted in turn
=====================  =====================================================

The original file is what gets hashed, stored and cited; the converted copy
is only read. A format nothing can read is reported, not dropped silently.
"""

from __future__ import annotations

import re
import shutil
import subprocess
import zipfile
from dataclasses import dataclass, field
from html import unescape
from pathlib import Path

#: Formats the extractors read directly.
NATIVE = {".csv", ".tsv", ".txt", ".xlsx", ".xlsm", ".pdf", ".docx"}
SPREADSHEETS = {".numbers", ".ods", ".xls"}
DOCUMENTS = {".doc", ".rtf", ".odt", ".html", ".htm"}
IMAGES = {".jpg", ".jpeg", ".png", ".tif", ".tiff", ".heic", ".heif", ".webp", ".bmp", ".gif"}
ARCHIVES = {".zip"}
ACCEPTED = NATIVE | SPREADSHEETS | DOCUMENTS | IMAGES | ARCHIVES

#: Don't unpack an archive bigger than this, or one nested deeper.
MAX_ARCHIVE_BYTES = 500 * 1024 * 1024
MAX_ARCHIVE_DEPTH = 2


class ConversionError(RuntimeError):
    pass


@dataclass
class Converted:
    """One readable file derived from an upload."""

    original: Path          # what was uploaded (hashed, stored, cited)
    readable: Path          # what the extractors open
    display_name: str       # name shown and cited, e.g. "production.zip › permit.pdf"
    notes: list[str] = field(default_factory=list)


def expand(path: Path, workdir: Path, *, depth: int = 0, prefix: str = "") -> list[Converted]:
    """All readable files an upload contains: itself, or an archive's members."""
    suffix = path.suffix.lower()
    name = f"{prefix}{path.name}"
    if suffix in ARCHIVES:
        if depth >= MAX_ARCHIVE_DEPTH:
            raise ConversionError(f"{name}: archive nested too deeply")
        out: list[Converted] = []
        target = workdir / f"unzipped-{depth}-{re.sub(r'[^A-Za-z0-9]+', '_', path.stem)}"
        target.mkdir(parents=True, exist_ok=True)
        with zipfile.ZipFile(path) as archive:
            total = sum(i.file_size for i in archive.infolist())
            if total > MAX_ARCHIVE_BYTES:
                raise ConversionError(f"{name}: archive expands to more than 500 MB")
            for info in archive.infolist():
                member = Path(info.filename)
                hidden = member.name.startswith((".", "~$")) or "__MACOSX" in member.parts
                if info.is_dir() or hidden:
                    continue
                # Flatten paths so a hostile archive cannot write outside target.
                dest = target / re.sub(r"[^A-Za-z0-9._ -]+", "_", member.name)
                with archive.open(info) as src, dest.open("wb") as dst:
                    shutil.copyfileobj(src, dst)
                out.extend(expand(dest, workdir, depth=depth + 1, prefix=f"{name} › "))
        return out
    return [convert(path, workdir, display_name=name)]


def convert(path: Path, workdir: Path, *, display_name: str | None = None) -> Converted:
    suffix = path.suffix.lower()
    name = display_name or path.name
    if suffix in NATIVE:
        return Converted(path, path, name)
    if suffix not in ACCEPTED:
        raise ConversionError(
            f"{name}: {suffix or 'files without an extension'} cannot be read. "
            f"Accepted: {', '.join(sorted(ACCEPTED))}"
        )
    workdir.mkdir(parents=True, exist_ok=True)
    out_stem = workdir / re.sub(r"[^A-Za-z0-9._-]+", "_", path.stem)
    if suffix in SPREADSHEETS:
        target = out_stem.with_suffix(".xlsx")
        tables = _read_spreadsheet(path)
        _write_xlsx(tables, target)
        return Converted(path, target, name,
                         [f"converted from {suffix} ({len(tables)} table(s)) for reading"])
    if suffix in DOCUMENTS:
        target = out_stem.with_suffix(".txt")
        target.write_text(_read_document(path), encoding="utf-8")
        return Converted(path, target, name, [f"converted from {suffix} to text for reading"])
    target = out_stem.with_suffix(".pdf")
    pages = _images_to_pdf(path, target)
    return Converted(path, target, name,
                     [f"photo/scan ({suffix}, {pages} page(s)) will be read by OCR"])


# --- spreadsheets -----------------------------------------------------------

def _cell(value) -> str:
    if value is None:
        return ""
    if hasattr(value, "isoformat"):
        return value.isoformat()[:10] if hasattr(value, "hour") and not value.hour \
            and not value.minute else value.isoformat()
    text = str(value).strip()
    return text[:-2] if re.fullmatch(r"-?\d+\.0", text) else text


def _read_spreadsheet(path: Path) -> list[tuple[str, list[list[str]]]]:
    suffix = path.suffix.lower()
    tables: list[tuple[str, list[list[str]]]] = []
    if suffix == ".numbers":
        from numbers_parser import Document

        doc = Document(str(path))
        for sheet in doc.sheets:
            for table in sheet.tables:
                rows = [[_cell(c.value) for c in row] for row in table.iter_rows()]
                tables.append((f"{sheet.name} - {table.name}"[:31], rows))
    elif suffix == ".xls":
        import xlrd

        book = xlrd.open_workbook(str(path))
        for sheet in book.sheets():
            rows = []
            for r in range(sheet.nrows):
                row = []
                for c in range(sheet.ncols):
                    cell = sheet.cell(r, c)
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        row.append(xlrd.xldate_as_datetime(cell.value, book.datemode)
                                   .date().isoformat())
                    else:
                        row.append(_cell(cell.value))
                rows.append(row)
            tables.append((sheet.name[:31], rows))
    elif suffix == ".ods":
        from odf import teletype
        from odf.opendocument import load
        from odf.table import Table, TableCell, TableRow

        doc = load(str(path))
        for table in doc.spreadsheet.getElementsByType(Table):
            rows = []
            for tr in table.getElementsByType(TableRow):
                row = []
                for tc in tr.getElementsByType(TableCell):
                    repeat = int(tc.getAttribute("numbercolumnsrepeated") or 1)
                    row.extend([teletype.extractText(tc).strip()] * min(repeat, 50))
                while row and not row[-1]:
                    row.pop()
                rows.append(row)
            tables.append(((table.getAttribute("name") or "Sheet")[:31], rows))
    if not any(any(any(c for c in r) for r in rows) for _, rows in tables):
        raise ConversionError(f"{path.name}: the spreadsheet has no data")
    return tables


def _write_xlsx(tables: list[tuple[str, list[list[str]]]], target: Path) -> None:
    from openpyxl import Workbook

    book = Workbook()
    book.remove(book.active)
    used: set[str] = set()
    for title, rows in tables:
        title = re.sub(r"[\[\]:*?/\\]", "-", title) or "Sheet"
        base, n = title, 2
        while title in used:
            title, n = f"{base[:28]}-{n}", n + 1
        used.add(title)
        sheet = book.create_sheet(title)
        for row in rows:
            sheet.append(row)
    book.save(target)


# --- documents --------------------------------------------------------------

def _read_document(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".doc":
        for tool in (["antiword", "-w", "0"], ["catdoc", "-w"]):
            if shutil.which(tool[0]):
                result = subprocess.run([*tool, str(path)], capture_output=True, timeout=120)
                if result.returncode == 0 and result.stdout.strip():
                    return result.stdout.decode("utf-8", errors="replace")
        raise ConversionError(f"{path.name}: no reader for old Word files is installed")
    if suffix == ".rtf":
        from striprtf.striprtf import rtf_to_text

        return rtf_to_text(path.read_text(encoding="latin-1", errors="replace"))
    if suffix == ".odt":
        from odf import teletype
        from odf.opendocument import load
        from odf.text import P

        doc = load(str(path))
        return "\n".join(teletype.extractText(p) for p in doc.getElementsByType(P))
    html = path.read_text(encoding="utf-8", errors="replace")
    html = re.sub(r"(?is)<(script|style)[^>]*>.*?</\1>", " ", html)
    html = re.sub(r"(?i)<br\s*/?>|</(p|div|tr|li|h\d)>", "\n", html)
    html = re.sub(r"(?i)</t[dh]>", "  |  ", html)
    return unescape(re.sub(r"<[^>]+>", " ", html))


# --- images -----------------------------------------------------------------

def _images_to_pdf(path: Path, target: Path) -> int:
    from PIL import Image, ImageOps

    if path.suffix.lower() in (".heic", ".heif"):
        from pillow_heif import register_heif_opener

        register_heif_opener()
    image = Image.open(path)
    frames = []
    for i in range(getattr(image, "n_frames", 1)):
        image.seek(i)
        # Phone photos carry their rotation in metadata; apply it, or OCR
        # reads the page sideways.
        frames.append(ImageOps.exif_transpose(image.copy()).convert("RGB"))
    frames[0].save(target, "PDF", resolution=300.0, save_all=True, append_images=frames[1:])
    return len(frames)
