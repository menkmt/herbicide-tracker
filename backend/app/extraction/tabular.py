"""Reading county "use record" exports: TSV, CSV and Excel.

This is the highest-volume source — a county hands over one file with every
reported application for a period.  Each *row* is a product line, and rows that
share a document number are one application, so the profile groups rows rather
than treating each as a record.

Nothing here assumes a fixed column order or fixed column names; columns are
located through :mod:`app.extraction.fieldmap`, and any column that is not
recognised is still preserved on the record's ``raw`` dictionary.
"""

from __future__ import annotations

import csv
import io
import re
from collections import OrderedDict
from pathlib import Path
from typing import Any

from app.core.confidence import Confidence
from app.core.provenance import ExtractionMethod, Provenance, SourceType, file_provenance
from app.core.siteid import SiteIdDecodeError, cross_check, decode_structured, try_decode
from app.extraction.base import (
    DataIssue,
    ExtractionResult,
    IssueCode,
    ProductApplication,
    PurRecord,
    normalize_method,
)
from app.extraction.fieldmap import (
    map_headers,
    match_label,
    parse_bool,
    parse_date,
    parse_datetime,
    parse_float,
    split_amount_units,
    split_commodity,
)

PROFILE_NAME = "tabular_use_records"

#: How many leading lines to inspect when hunting for the header row.  County
#: exports often carry a title and a blank line above the real header.
HEADER_SEARCH_DEPTH = 15

#: A row must match at least this many known labels to be the header.
MIN_HEADER_MATCHES = 4

#: Fields that describe the application as a whole rather than one product.
RECORD_FIELDS = (
    "document_number", "permit_number", "county_name", "county_code", "site_district",
    "operator_name", "applicator_name", "applicator_license", "applicator_license_type",
    "applicator_address", "location_text", "site_id", "mtrs_text", "commodity",
    "commodity_code", "submittal_status", "school_notification", "pca_name",
)


def _read_delimited(path: Path) -> list[list[str]]:
    text = path.read_text(encoding="utf-8", errors="replace")
    sample = "\n".join(text.splitlines()[:40])
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters="\t,;|")
        delimiter = dialect.delimiter
    except csv.Error:
        # Sniffing fails on files with a prose title line; fall back to
        # whichever candidate appears most often.
        delimiter = max("\t,;|", key=sample.count)
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    return [row for row in reader]


def _read_excel(path: Path) -> list[list[str]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows: list[list[str]] = []
    for row in sheet.iter_rows(values_only=True):
        rows.append(["" if cell is None else str(cell).strip() for cell in row])
    workbook.close()
    return rows


def read_rows(path: str | Path) -> list[list[str]]:
    path = Path(path)
    if path.suffix.lower() in {".xlsx", ".xlsm", ".xls"}:
        return _read_excel(path)
    return _read_delimited(path)


def find_header_row(rows: list[list[str]]) -> tuple[int, dict[int, str]] | None:
    """Locate the header row and its column mapping.

    Returns ``None`` when no row looks like a header, which is how the registry
    decides a file is not a tabular use-record export.
    """
    best: tuple[int, dict[int, str]] | None = None
    best_score = 0
    for index, row in enumerate(rows[:HEADER_SEARCH_DEPTH]):
        mapping = map_headers([str(c) for c in row])
        if len(mapping) > best_score:
            best_score = len(mapping)
            best = (index, mapping)
    if best is None or best_score < MIN_HEADER_MATCHES:
        return None
    return best


def _row_dict(row: list[str], headers: list[str]) -> dict[str, str]:
    return {
        headers[i] if i < len(headers) else f"column_{i}": (row[i] if i < len(row) else "")
        for i in range(max(len(row), len(headers)))
    }


def _get(values: dict[str, str], field: str) -> str | None:
    value = values.get(field)
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _resolve_site(
    record: PurRecord,
    values: dict[str, str],
    *,
    county: str | None,
    provenance: Provenance,
) -> None:
    """Work out where the application happened, preferring explicit columns.

    Order matters.  Separate township/range/section columns carry the direction
    letters explicitly, so they beat re-parsing a packed site ID that has to
    infer them.  Where a site ID is *also* present it is used to corroborate.
    Site IDs are not always a packed PLSS code — some counties use a name such
    as ``TBarr`` — so a site ID that will not decode is only a problem when
    there is nothing else to fall back on.
    """
    township = _get(values, "township")
    range_ = _get(values, "range")
    section = _get(values, "section")
    mtrs_text = _get(values, "mtrs_text")
    site_id = record.site_id

    if township and range_ and section:
        # Direction letters are usually fused to the number ("29N", "11E").
        t_match = re.match(r"(\d+)\s*([NS])?", township.upper())
        r_match = re.match(r"(\d+)\s*([EW])?", range_.upper())
        try:
            record.site = decode_structured(
                township=t_match.group(1) if t_match else township,
                township_dir=(t_match.group(2) if t_match else None) or _get(values, "township_dir"),
                range_=r_match.group(1) if r_match else range_,
                range_dir=(r_match.group(2) if r_match else None) or _get(values, "range_dir"),
                section=section,
                meridian=_get(values, "meridian"),
                county=county,
                raw=site_id or f"{township}{range_}{section}",
            )
            record.field_sources["site"] = provenance.at(
                "township/range/section columns"
            ).with_confidence(record.site.confidence)
        except SiteIdDecodeError as exc:
            record.add_issue(
                IssueCode.UNPARSED_SECTION,
                f"township/range/section columns could not be read: {exc}",
                field_name="site",
            )

        # Corroborate against the packed site ID when there is one.
        if record.site and site_id:
            decoded_from_id = try_decode(site_id, county=county)
            if decoded_from_id and decoded_from_id.trs != record.site.trs:
                record.add_issue(
                    IssueCode.SITE_ID_MTRS_CONFLICT,
                    f"site ID {site_id} decodes to {decoded_from_id.trs} but the "
                    f"township/range/section columns say {record.site.trs}",
                    field_name="site",
                )
        return

    if site_id or mtrs_text:
        try:
            checked = cross_check(site_id, mtrs_text, county=county)
        except SiteIdDecodeError:
            record.add_issue(
                IssueCode.UNDECODABLE_SITE_ID,
                f"site ID {site_id!r} is not a PLSS code and no township/range/section "
                "columns were supplied, so the application cannot be located",
                field_name="site_id",
            )
            return
        record.site = checked.decoded
        record.field_sources["site"] = provenance.at(
            f"site ID {site_id}" if site_id else f"MTRS {mtrs_text}"
        ).with_confidence(checked.decoded.confidence)
        if not checked.agrees and checked.from_site_id and checked.from_mtrs:
            record.add_issue(IssueCode.SITE_ID_MTRS_CONFLICT, checked.detail, field_name="site")
        return

    record.add_issue(
        IssueCode.MISSING_LOCATION,
        "no site ID, MTRS or township/range/section columns were present",
        field_name="site",
    )


def _build_product(values: dict[str, str], line_number: int) -> ProductApplication | None:
    name = _get(values, "product_name")
    reg = _get(values, "epa_reg_no")
    if not name and not reg:
        return None
    treated_amount, treated_units = split_amount_units(_get(values, "treated_amount"))
    return ProductApplication(
        product_name=name,
        epa_reg_no=reg,
        quantity=parse_float(_get(values, "quantity")),
        quantity_units=_get(values, "quantity_units"),
        treated_amount=treated_amount,
        treated_units=_get(values, "treated_units") or treated_units,
        registration_expired=parse_bool(_get(values, "registration_expired")),
        line_number=line_number,
    )


def _grouping_key(values: dict[str, str]) -> tuple:
    """Rows sharing this key are one application.

    The document number is the real identifier.  When a county export omits it,
    fall back to the combination that uniquely identifies a use report.
    """
    document = _get(values, "document_number")
    if document:
        return ("doc", document)
    return (
        "composite",
        _get(values, "permit_number") or "",
        _get(values, "site_id") or "",
        _get(values, "application_date") or _get(values, "start_datetime") or "",
        _get(values, "application_time") or "",
    )


def extract(
    path: str | Path,
    *,
    county: str | None = None,
    sha256: str | None = None,
    source_name: str | None = None,
) -> ExtractionResult:
    """Extract PUR records from a delimited or Excel use-record export."""
    path = Path(path)
    name = source_name or path.name
    result = ExtractionResult(source_name=name, sha256=sha256, profile=PROFILE_NAME)

    rows = read_rows(path)
    header = find_header_row(rows)
    if header is None:
        result.issues.append(
            DataIssue(
                "unrecognised_format",
                "no row in this file matched enough known PUR column names to be a header",
                severity="review",
            )
        )
        return result

    header_index, mapping = header
    headers = [mapping.get(i, f"column_{i}") for i in range(len(rows[header_index]))]
    original_headers = [str(c) for c in rows[header_index]]
    unmapped = [h for i, h in enumerate(original_headers) if i not in mapping and h.strip()]
    if unmapped:
        result.notes.append(f"columns kept but not mapped: {', '.join(unmapped)}")
    result.notes.append(f"header found on line {header_index + 1}; {len(mapping)} columns mapped")

    base_provenance = file_provenance(
        source_type=SourceType.PUR_USE_RECORD,
        file_name=name,
        sha256=sha256,
        method=ExtractionMethod.TABULAR_COLUMN,
        confidence=Confidence.VERIFIED,
    )

    groups: OrderedDict[tuple, list[tuple[int, dict[str, str]]]] = OrderedDict()
    for offset, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        if not any(str(c).strip() for c in row):
            continue
        values = _row_dict([str(c) for c in row], headers)
        groups.setdefault(_grouping_key(values), []).append((offset, values))

    for key, members in groups.items():
        line_numbers, value_rows = zip(*members, strict=True)
        first = value_rows[0]
        record = PurRecord(source_profile=PROFILE_NAME)
        record.raw = {
            "rows": [dict(v) for v in value_rows],
            "source_lines": list(line_numbers),
            "grouping_key": list(key),
        }

        provenance = base_provenance.at(f"line {line_numbers[0]}")
        if key[0] != "doc":
            record.add_issue(
                "no_document_number",
                "this export has no document number column; rows were grouped by "
                "permit, site and date instead",
                severity="note",
            )

        for field_name in RECORD_FIELDS:
            value = _get(first, field_name)
            if value is None:
                continue
            setattr(record, field_name, value)
            record.field_sources[field_name] = provenance.at(
                f"line {line_numbers[0]}, column {field_name!r}"
            )

        if record.commodity and not record.commodity_code:
            commodity, code = split_commodity(record.commodity)
            record.commodity, record.commodity_code = commodity, code

        record.county_name = record.county_name or county

        # --- when ---------------------------------------------------------
        start = parse_datetime(
            _get(first, "start_datetime") or _get(first, "application_date"),
            time_value=_get(first, "application_time"),
        )
        end = parse_datetime(_get(first, "end_datetime")) or None
        record.start_datetime = start
        record.end_datetime = end
        record.application_date = (
            start.date() if start else parse_date(_get(first, "application_date"))
        )
        if record.application_date is None:
            record.add_issue(
                IssueCode.MISSING_DATE,
                "no readable application date",
                field_name="application_date",
            )
        else:
            record.field_sources["application_date"] = provenance.at(
                "application date column"
            )

        # --- how ----------------------------------------------------------
        record.method_raw = _get(first, "method")
        record.method = normalize_method(record.method_raw)
        if record.method_raw:
            record.field_sources["method"] = provenance.at("application method column")

        planted_amount, planted_units = split_amount_units(_get(first, "planted_amount"))
        record.planted_amount = planted_amount
        record.planted_units = _get(first, "planted_units") or planted_units
        treated_amount, treated_units = split_amount_units(_get(first, "treated_amount"))
        record.treated_amount = treated_amount
        record.treated_units = _get(first, "treated_units") or treated_units

        # --- where --------------------------------------------------------
        _resolve_site(record, first, county=record.county_name, provenance=provenance)

        # --- what ---------------------------------------------------------
        for line_number, values in members:
            product = _build_product(values, line_number)
            if product is not None:
                record.products.append(product)
        if not record.products:
            record.add_issue(
                IssueCode.UNKNOWN_PRODUCT,
                "no product name or registration number on any row of this record",
            )
        # Treated acreage is usually repeated on every product row; lift it to
        # the record when the rows agree, so the public row can show one figure.
        if record.treated_amount is None:
            amounts = {p.treated_amount for p in record.products if p.treated_amount is not None}
            if len(amounts) == 1:
                record.treated_amount = amounts.pop()
                record.treated_units = next(
                    (p.treated_units for p in record.products if p.treated_units), None
                )

        if not record.operator_name:
            record.add_issue(
                IssueCode.MISSING_OWNER,
                "no permittee or operator name on this record",
                field_name="operator_name",
            )

        result.records.append(record)

    return result


def sniff(path: str | Path) -> float:
    """Confidence (0-1) that this file is a tabular use-record export."""
    path = Path(path)
    if path.suffix.lower() not in {".csv", ".tsv", ".txt", ".xlsx", ".xlsm", ".xls"}:
        return 0.0
    try:
        rows = read_rows(path)
    except Exception:
        return 0.0
    header = find_header_row(rows)
    if header is None:
        return 0.0
    _, mapping = header
    # More recognised columns means more certainty; 12+ is a full export.
    return min(1.0, 0.4 + 0.05 * len(mapping))


def _unused(*args: Any) -> None:  # pragma: no cover - keeps linters honest
    del args
